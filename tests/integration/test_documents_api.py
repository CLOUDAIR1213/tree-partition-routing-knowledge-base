import logging
from io import BytesIO
from sqlite3 import connect

from docx import Document
from reportlab.pdfgen import canvas


def upload_markdown(client, content: bytes | None = None, filename: str = "policy.md"):
    return client.post(
        "/api/v1/documents",
        files={"file": (filename, content or b"# Policy\n\nDemo amount is 600 yuan.", "text/markdown")},
        data={"partition": "finance", "title": "Demo policy"},
    )


def test_upload_preview_and_approve_into_confirmed_partition(
    client, fake_registry, caplog
):
    caplog.set_level(logging.INFO, logger="app.services.review")
    upload = upload_markdown(client)
    assert upload.status_code == 201, upload.text
    payload = upload.json()
    assert payload["status"] == "pending_review"
    assert payload["confirmed_partition"] is None
    assert upload.headers["X-Request-ID"] == payload["request_id"]

    document_id = payload["document_id"]
    detail = client.get(f"/api/v1/documents/{document_id}")
    assert detail.status_code == 200
    assert "stored_path" not in detail.json()
    assert "checksum_sha256" not in detail.json()
    assert detail.json()["title"] == "Demo policy"
    assert detail.json()["parse_quality"] == payload["parse_quality"]

    preview = client.get(f"/api/v1/documents/{document_id}/preview?limit=1&offset=0")
    assert preview.status_code == 200
    assert preview.json()["total"] == payload["chunk_count"]
    assert "embedding_text" not in preview.json()["items"][0]
    assert preview.json()["items"][0]["text"] == "Demo amount is 600 yuan."
    assert preview.json()["items"][0]["title"] == "Demo policy"
    assert preview.json()["parse_quality"] == payload["parse_quality"]
    assert fake_registry.upsert_calls == []

    review = client.post(
        f"/api/v1/documents/{document_id}/review",
        json={
            "action": "approve",
            "confirmed_partition": "tech",
            "reviewer_name": "demo-reviewer",
            "note": "Corrected partition",
        },
    )
    assert review.status_code == 200, review.text
    assert review.json()["status"] == "ready"
    assert review.json()["selected_partition"] == "finance"
    assert review.json()["confirmed_partition"] == "tech"
    assert fake_registry.upsert_calls == ["tech"]
    assert fake_registry.rows["finance"] == {}
    assert fake_registry.rows["hr"] == {}
    messages = [record.getMessage() for record in caplog.records]
    assert any("level=chunk operation=upsert" in message for message in messages)
    assert any("level=document operation=upsert" in message for message in messages)
    assert any("level=section operation=upsert" in message for message in messages)
    assert any("review_index_completed" in message for message in messages)


def test_upload_without_title_uses_filename_stem_not_first_section_heading(client):
    upload = client.post(
        "/api/v1/documents",
        files={
            "file": (
                "company-security-handbook.md",
                b"# Document control information\n\nSecurity policy content.",
                "text/markdown",
            )
        },
        data={"partition": "tech"},
    )
    assert upload.status_code == 201, upload.text

    document_id = upload.json()["document_id"]
    detail = client.get(f"/api/v1/documents/{document_id}")
    preview = client.get(f"/api/v1/documents/{document_id}/preview")

    assert detail.status_code == 200
    assert detail.json()["title"] == "company-security-handbook"
    assert preview.status_code == 200
    assert preview.json()["items"][0]["title"] == "company-security-handbook"


def test_reject_does_not_touch_any_index(client, fake_registry):
    document_id = upload_markdown(client, b"# Leave\n\nDemo leave policy.", "leave.md").json()["document_id"]
    response = client.post(
        f"/api/v1/documents/{document_id}/review",
        json={"action": "reject", "confirmed_partition": None, "reviewer_name": None, "note": "Incomplete"},
    )
    assert response.status_code == 200
    assert response.json()["status"] == "rejected"
    assert fake_registry.upsert_calls == []


def test_review_action_validation_uses_stable_error_codes(client):
    first = upload_markdown(client, b"# First\n\nReview validation one.", "first.md").json()["document_id"]
    missing_partition = client.post(
        f"/api/v1/documents/{first}/review",
        json={"action": "approve", "confirmed_partition": None, "reviewer_name": None, "note": None},
    )
    assert missing_partition.status_code == 422
    assert missing_partition.json()["code"] == "REVIEW_PARTITION_REQUIRED"

    second = upload_markdown(client, b"# Second\n\nReview validation two.", "second.md").json()["document_id"]
    reject_with_partition = client.post(
        f"/api/v1/documents/{second}/review",
        json={"action": "reject", "confirmed_partition": "hr", "reviewer_name": None, "note": None},
    )
    assert reject_with_partition.status_code == 422
    assert reject_with_partition.json()["code"] == "VALIDATION_ERROR"


def test_duplicate_upload_is_rejected_with_existing_document(client):
    content = b"# Stable\n\nSame content."
    first = upload_markdown(client, content)
    second = upload_markdown(client, content)
    assert first.status_code == 201
    assert second.status_code == 409
    assert second.json()["code"] == "DUPLICATE_DOCUMENT"
    assert second.json()["details"]["document_id"] == first.json()["document_id"]


def test_errors_use_contract_and_request_id(client):
    response = client.post(
        "/api/v1/documents",
        files={"file": ("policy.md", b"content", "text/markdown")},
        data={"partition": "Finance"},
        headers={"X-Request-ID": "req_frontend_123"},
    )
    assert response.status_code == 422
    assert response.json()["code"] == "INVALID_PARTITION"
    assert response.json()["request_id"] == "req_frontend_123"
    assert response.headers["X-Request-ID"] == "req_frontend_123"
    assert "detail" not in response.json()


def test_path_traversal_filename_cannot_escape_storage(client, tmp_path):
    response = upload_markdown(client, filename="../../outside.md")
    assert response.status_code == 201
    assert response.json()["original_filename"] == "outside.md"
    assert not (tmp_path / "outside.md").exists()


def test_list_pagination_and_health(client):
    upload_markdown(client, b"# A\n\nFirst demo.", "a.md")
    upload_markdown(client, b"# B\n\nSecond demo.", "b.md")
    listing = client.get("/api/v1/documents?status=pending_review&limit=1&offset=1")
    assert listing.status_code == 200
    assert listing.json()["total"] == 2
    assert listing.json()["limit"] == 1
    assert listing.json()["offset"] == 1
    health = client.get("/api/v1/health")
    assert health.status_code == 200
    assert health.json()["indexes"] == {
        partition: {"document": "ready", "section": "ready", "chunk": "ready"}
        for partition in ("finance", "hr", "tech")
    }
    assert health.json()["router"] == "not_configured"
    assert health.json()["answer"] == "not_configured"
    assert health.json()["web_search"] == "disabled"


def test_list_filters_by_effective_partition_and_query(client):
    finance_id = upload_markdown(
        client,
        b"# Travel handbook\n\nExpense submission rules.",
        "travel-handbook.md",
    ).json()["document_id"]
    tech_id = upload_markdown(
        client,
        b"# VPN handbook\n\nCertificate troubleshooting.",
        "vpn-handbook.md",
    ).json()["document_id"]
    review = client.post(
        f"/api/v1/documents/{tech_id}/review",
        json={
            "action": "approve",
            "confirmed_partition": "tech",
            "reviewer_name": None,
            "note": None,
        },
    )
    assert review.status_code == 200

    by_query = client.get("/api/v1/documents?q=vPn-HaNdBoOk")
    assert by_query.status_code == 200
    assert [item["document_id"] for item in by_query.json()["items"]] == [tech_id]

    by_effective_partition = client.get("/api/v1/documents?partition=tech")
    assert by_effective_partition.status_code == 200
    assert [item["document_id"] for item in by_effective_partition.json()["items"]] == [
        tech_id
    ]

    pending_finance = client.get(
        "/api/v1/documents?partition=finance&status=pending_review"
    )
    assert pending_finance.status_code == 200
    assert [
        item["document_id"] for item in pending_finance.json()["items"]
    ] == [finance_id]


def test_list_rejects_invalid_partition_and_long_query(client):
    invalid_partition = client.get("/api/v1/documents?partition=legal")
    assert invalid_partition.status_code == 422
    assert invalid_partition.json()["code"] == "INVALID_PARTITION"

    long_query = client.get(f"/api/v1/documents?q={'x' * 201}")
    assert long_query.status_code == 422
    assert long_query.json()["code"] == "VALIDATION_ERROR"


def test_framework_404_uses_error_contract(client):
    response = client.get("/api/v1/not-a-route")
    assert response.status_code == 404
    assert response.json()["code"] == "NOT_FOUND"
    assert response.headers["X-Request-ID"] == response.json()["request_id"]


def test_txt_pdf_docx_and_markdown_uploads(client):
    docx_buffer = BytesIO()
    document = Document()
    document.add_heading("DOCX Demo", level=1)
    document.add_paragraph("Approval material list.")
    document.save(docx_buffer)

    pdf_buffer = BytesIO()
    pdf = canvas.Canvas(pdf_buffer)
    pdf.drawString(72, 720, "PDF demo approval process")
    pdf.save()

    samples = [
        ("demo.txt", b"Plain text demo policy.", "text/plain"),
        ("demo.md", b"# Markdown\n\nMarkdown demo policy.", "text/markdown"),
        (
            "demo.docx",
            docx_buffer.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
        ("demo.pdf", pdf_buffer.getvalue(), "application/pdf"),
    ]
    for filename, content, mime_type in samples:
        response = client.post(
            "/api/v1/documents",
            files={"file": (filename, content, mime_type)},
            data={"partition": "hr"},
        )
        assert response.status_code == 201, response.text
        assert response.json()["status"] == "pending_review"
        assert response.json()["chunk_count"] >= 1
        assert response.json()["parse_quality"]["source_format"] in {
            "text",
            "markdown",
            "docx",
            "pdf",
        }


def test_upload_returns_quality_and_non_authoritative_partition_suggestion(client):
    response = upload_markdown(
        client,
        (
            "# \u5dee\u65c5\u62a5\u9500\n\n"
            "\u5dee\u65c5\u8d39\u7528\u62a5\u9500\u5236\u5ea6\u3002\n\n"
            "## \u53d1\u7968\u6750\u6599\n\n"
            "\u62a5\u9500\u7533\u8bf7\u9700\u9644\u53d1\u7968\u3001\u5dee\u65c5\u5ba1\u6279\u5355\u548c\u8d39\u7528\u660e\u7ec6\u3002"
        ).encode(),
        "travel.md",
    )

    assert response.status_code == 201, response.text
    quality = response.json()["parse_quality"]
    assert quality["section_count"] == 2
    assert quality["heading_recognition_rate"] == 1
    assert quality["partition_suggestion"]["partition"] == "finance"
    assert quality["partition_suggestion"]["confidence"] > 0
    assert response.json()["selected_partition"] == "finance"
    assert response.json()["confirmed_partition"] is None


def test_content_and_extension_mismatch_is_rejected(client):
    response = client.post(
        "/api/v1/documents",
        files={"file": ("fake.pdf", b"not a pdf", "application/pdf")},
        data={"partition": "finance"},
    )
    assert response.status_code == 415
    assert response.json()["code"] == "UNSUPPORTED_FILE_TYPE"


def test_index_failure_compensates_and_marks_document_failed(client, fake_registry):
    document_id = upload_markdown(
        client, b"# Failure\n\nCompensation path unique content.", "failure.md"
    ).json()["document_id"]
    fake_registry.fail_after_upsert = True
    response = client.post(
        f"/api/v1/documents/{document_id}/review",
        json={
            "action": "approve",
            "confirmed_partition": "finance",
            "reviewer_name": None,
            "note": None,
        },
    )
    assert response.status_code == 500
    assert response.json()["code"] == "INDEX_WRITE_FAILED"
    assert fake_registry.delete_calls == ["finance"]
    assert fake_registry.rows["finance"] == {}
    detail = client.get(f"/api/v1/documents/{document_id}")
    assert detail.json()["status"] == "failed"
    assert detail.json()["error_code"] == "INDEX_WRITE_FAILED"


def test_delete_pending_document_removes_files_chunks_and_metadata(client, fake_registry):
    document_id = upload_markdown(
        client,
        b"# Disposable\n\nPending content to remove.",
        "disposable.md",
    ).json()["document_id"]
    settings = client.app.state.settings
    assert (settings.raw_root / document_id).exists()
    assert (settings.staging_root / document_id).exists()

    response = client.delete(f"/api/v1/documents/{document_id}")

    assert response.status_code == 200, response.text
    assert response.json()["deleted"] is True
    assert response.json()["removed_chunk_count"] == 1
    assert fake_registry.delete_calls == ["finance", "hr", "tech"]
    assert not (settings.raw_root / document_id).exists()
    assert not (settings.staging_root / document_id).exists()
    assert client.get(f"/api/v1/documents/{document_id}").status_code == 404
    assert client.get(f"/api/v1/documents/{document_id}/preview").status_code == 404


def test_delete_ready_document_removes_indexed_chunks_from_every_partition(
    client,
    fake_registry,
):
    document_id = upload_markdown(
        client,
        b"# Indexed delete\n\nRemove indexed content too.",
        "indexed-delete.md",
    ).json()["document_id"]
    assert client.post(
        f"/api/v1/documents/{document_id}/review",
        json={"action": "approve", "confirmed_partition": "tech"},
    ).status_code == 200
    assert fake_registry.rows["tech"]

    response = client.delete(f"/api/v1/documents/{document_id}")

    assert response.status_code == 200, response.text
    assert all(not rows for rows in fake_registry.rows.values())
    assert client.get(f"/api/v1/documents/{document_id}").status_code == 404


def test_change_partition_moves_every_chunk_and_updates_review_fields(
    client,
    fake_registry,
):
    document_id = upload_markdown(
        client,
        b"# Reassign\n\nMove this indexed content.",
        "reassign.md",
    ).json()["document_id"]
    approved = client.post(
        f"/api/v1/documents/{document_id}/review",
        json={
            "action": "approve",
            "confirmed_partition": "finance",
            "reviewer_name": "first-reviewer",
            "note": "Initial partition",
        },
    )
    assert approved.status_code == 200
    chunk_ids = set(fake_registry.rows["finance"])

    response = client.post(
        f"/api/v1/documents/{document_id}/partition",
        json={
            "confirmed_partition": "tech",
            "reviewer_name": "partition-reviewer",
            "note": "Technical ownership confirmed",
        },
    )

    assert response.status_code == 200, response.text
    assert response.json()["previous_partition"] == "finance"
    assert response.json()["confirmed_partition"] == "tech"
    assert response.json()["status"] == "ready"
    assert set(fake_registry.rows["tech"]) == chunk_ids
    assert fake_registry.rows["finance"] == {}
    assert fake_registry.rows["hr"] == {}
    detail = client.get(f"/api/v1/documents/{document_id}").json()
    assert detail["confirmed_partition"] == "tech"
    assert detail["review_note"] == "Technical ownership confirmed"


def test_change_partition_failure_restores_original_ready_index(client, fake_registry):
    document_id = upload_markdown(
        client,
        b"# Restore\n\nKeep the original partition on failure.",
        "restore.md",
    ).json()["document_id"]
    assert client.post(
        f"/api/v1/documents/{document_id}/review",
        json={"action": "approve", "confirmed_partition": "finance"},
    ).status_code == 200
    original_ids = set(fake_registry.rows["finance"])
    fake_registry.fail_next_upsert = True

    response = client.post(
        f"/api/v1/documents/{document_id}/partition",
        json={"confirmed_partition": "hr"},
    )

    assert response.status_code == 500
    assert response.json()["code"] == "PARTITION_CHANGE_FAILED"
    assert response.json()["details"]["compensation"] == "compensation_succeeded"
    assert set(fake_registry.rows["finance"]) == original_ids
    assert fake_registry.rows["hr"] == {}
    detail = client.get(f"/api/v1/documents/{document_id}").json()
    assert detail["status"] == "ready"
    assert detail["confirmed_partition"] == "finance"


def test_reopen_review_removes_indexes_and_allows_new_approval(client, fake_registry):
    document_id = upload_markdown(
        client,
        b"# Review again\n\nRecheck this document.",
        "review-again.md",
    ).json()["document_id"]
    assert client.post(
        f"/api/v1/documents/{document_id}/review",
        json={"action": "approve", "confirmed_partition": "finance"},
    ).status_code == 200

    reopened = client.post(
        f"/api/v1/documents/{document_id}/reopen-review",
        json={"reviewer_name": None, "note": "Needs another review"},
    )

    assert reopened.status_code == 200, reopened.text
    assert reopened.json()["previous_partition"] == "finance"
    assert reopened.json()["confirmed_partition"] is None
    assert reopened.json()["status"] == "pending_review"
    assert all(not rows for rows in fake_registry.rows.values())
    detail = client.get(f"/api/v1/documents/{document_id}").json()
    assert detail["confirmed_partition"] is None
    assert detail["review_note"] == "Needs another review"

    approved = client.post(
        f"/api/v1/documents/{document_id}/review",
        json={"action": "approve", "confirmed_partition": "hr"},
    )
    assert approved.status_code == 200
    assert approved.json()["confirmed_partition"] == "hr"
    assert fake_registry.rows["finance"] == {}
    assert fake_registry.rows["hr"]


def test_management_actions_reject_invalid_states_and_unchanged_partition(client):
    document_id = upload_markdown(
        client,
        b"# State rules\n\nManagement state validation.",
        "state-rules.md",
    ).json()["document_id"]

    pending_change = client.post(
        f"/api/v1/documents/{document_id}/partition",
        json={"confirmed_partition": "tech"},
    )
    pending_reopen = client.post(
        f"/api/v1/documents/{document_id}/reopen-review",
        json={},
    )
    assert pending_change.status_code == 409
    assert pending_change.json()["code"] == "INVALID_DOCUMENT_STATE"
    assert pending_reopen.status_code == 409
    assert pending_reopen.json()["code"] == "INVALID_DOCUMENT_STATE"

    assert client.post(
        f"/api/v1/documents/{document_id}/review",
        json={"action": "approve", "confirmed_partition": "finance"},
    ).status_code == 200
    unchanged = client.post(
        f"/api/v1/documents/{document_id}/partition",
        json={"confirmed_partition": "finance"},
    )
    assert unchanged.status_code == 422
    assert unchanged.json()["code"] == "PARTITION_UNCHANGED"


def test_tree_nodes_follow_approval_partition_change_and_reopen(
    hierarchical_client,
    fake_registry,
    fake_hierarchy_registry,
):
    document_id = upload_markdown(
        hierarchical_client,
        b"# Runbook\n\nInitial technical process.\n\n## Recovery\n\nRecovery steps.",
        "tree-lifecycle.md",
    ).json()["document_id"]

    approved = hierarchical_client.post(
        f"/api/v1/documents/{document_id}/review",
        json={"action": "approve", "confirmed_partition": "finance"},
    )

    assert approved.status_code == 200, approved.text
    assert fake_registry.rows["finance"]
    assert fake_hierarchy_registry.rows[("finance", "document")]
    assert fake_hierarchy_registry.rows[("finance", "section")]
    database_path = (
        hierarchical_client.app.state.settings.data_root / "metadata" / "knowledge.db"
    )
    with connect(database_path) as database:
        persisted = database.execute(
            "SELECT parent_id, level, partition, document_id, indexed_at "
            "FROM hierarchy_nodes ORDER BY level, id"
        ).fetchall()
    document_nodes = [row for row in persisted if row[1] == "document"]
    section_nodes = [row for row in persisted if row[1] == "section"]
    assert len(document_nodes) == 1
    assert section_nodes
    assert document_nodes[0][0] is None
    assert document_nodes[0][1:4] == ("document", "finance", document_id)
    assert document_nodes[0][4] is not None
    assert all(row[0] == f"hdoc:{document_id}" for row in section_nodes)
    assert all(row[1:4] == ("section", "finance", document_id) for row in section_nodes)
    assert all(row[4] is not None for row in section_nodes)

    moved = hierarchical_client.post(
        f"/api/v1/documents/{document_id}/partition",
        json={"confirmed_partition": "tech"},
    )

    assert moved.status_code == 200, moved.text
    assert fake_hierarchy_registry.rows[("tech", "document")]
    assert fake_hierarchy_registry.rows[("tech", "section")]
    assert fake_hierarchy_registry.rows[("finance", "document")] == {}
    assert fake_hierarchy_registry.rows[("finance", "section")] == {}
    assert fake_hierarchy_registry.rows[("hr", "document")] == {}
    assert fake_hierarchy_registry.rows[("hr", "section")] == {}

    reopened = hierarchical_client.post(
        f"/api/v1/documents/{document_id}/reopen-review",
        json={"note": "Recheck tree"},
    )

    assert reopened.status_code == 200, reopened.text
    assert all(not rows for rows in fake_hierarchy_registry.rows.values())


def test_tree_nodes_are_removed_when_a_ready_document_is_deleted(
    hierarchical_client,
    fake_hierarchy_registry,
):
    document_id = upload_markdown(
        hierarchical_client,
        b"# Disposable\n\nRemove tree nodes with this document.",
        "tree-delete.md",
    ).json()["document_id"]
    assert hierarchical_client.post(
        f"/api/v1/documents/{document_id}/review",
        json={"action": "approve", "confirmed_partition": "tech"},
    ).status_code == 200

    deleted = hierarchical_client.delete(f"/api/v1/documents/{document_id}")

    assert deleted.status_code == 200, deleted.text
    assert all(not rows for rows in fake_hierarchy_registry.rows.values())


def test_tree_index_failure_compensates_chunk_and_node_writes(
    hierarchical_client,
    fake_registry,
    fake_hierarchy_registry,
):
    document_id = upload_markdown(
        hierarchical_client,
        b"# Tree failure\n\nCompensate every index surface.",
        "tree-failure.md",
    ).json()["document_id"]
    fake_hierarchy_registry.fail_next_node_upsert = True

    response = hierarchical_client.post(
        f"/api/v1/documents/{document_id}/review",
        json={"action": "approve", "confirmed_partition": "tech"},
    )

    assert response.status_code == 500
    assert response.json()["code"] == "INDEX_WRITE_FAILED"
    assert response.json()["details"]["compensation"] == "compensation_succeeded"
    assert fake_registry.rows["tech"] == {}
    assert all(not rows for rows in fake_hierarchy_registry.rows.values())
