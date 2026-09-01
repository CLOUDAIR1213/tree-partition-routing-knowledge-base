from pathlib import Path

from docx import Document
from reportlab.pdfgen import canvas

from app.services.chunker import SectionChunker
from app.services.parser import StandardDocumentParser


def test_markdown_sections_and_stable_chunk_ids(tmp_path: Path):
    path = tmp_path / "policy.md"
    path.write_text("# Policy\n\nIntro.\n\n## Limit\n\nAmount 600.\n", encoding="utf-8")
    parser = StandardDocumentParser()
    sections = parser.parse(path, "text/markdown")
    assert [section.section_path for section in sections] == ["Policy", "Policy > Limit"]
    chunker = SectionChunker(20, 2, 30, 5)
    first = chunker.chunk("doc_123", "Policy", sections)
    second = chunker.chunk("doc_123", "Policy", sections)
    assert [item.chunk_id for item in first] == [item.chunk_id for item in second]
    assert [item.checksum_sha256 for item in first] == [item.checksum_sha256 for item in second]


def test_docx_parser_preserves_heading_hierarchy(tmp_path: Path):
    path = tmp_path / "policy.docx"
    document = Document()
    document.add_heading("Handbook", level=1)
    document.add_paragraph("Introduction")
    document.add_heading("Approval", level=2)
    document.add_paragraph("Submit the demo form.")
    document.save(path)
    sections = StandardDocumentParser().parse(
        path, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert [section.section_path for section in sections] == ["Handbook", "Handbook > Approval"]


def test_pdf_parser_returns_page_metadata(tmp_path: Path):
    path = tmp_path / "policy.pdf"
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 720, "Demo travel limit: 600 yuan")
    pdf.showPage()
    pdf.drawString(72, 720, "Demo approval process")
    pdf.save()
    sections = StandardDocumentParser().parse(path, "application/pdf")
    assert len(sections) == 2
    assert sections[0].page_start == 1
    assert sections[1].page_end == 2


def test_long_english_window_preserves_word_boundaries():
    text = " ".join(f"word{index}" for index in range(80))
    chunker = SectionChunker(20, 5, 25, 5)
    parts = chunker._window_split(text)
    assert len(parts) > 1
    assert "word0 word1" in parts[0]
    assert "word0word1" not in parts[0]
